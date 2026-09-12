#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
国家医保智能监管知识库、规则库 入库脚本
用法: python3 kb_import_insurance.py <解压后的目录>  [--dry]

处理:
  1. 每个批次文件生成一份结构化 markdown -> medical/医保监管/
  2. 原始 xlsx/pdf 备份 -> source/国家医保智能监管知识库规则库/
  3. 知识点明细写入 SQLite insurance_rules、药品/项目代码写入 insurance_codes
  4. documents 表登记每个批次
"""
import os, re, sys, glob, json, shutil, sqlite3
from datetime import datetime
from openpyxl import load_workbook

KB_ROOT = os.path.expanduser('~/private_db/knowledge')
DB_PATH = os.path.join(KB_ROOT, 'db', 'knowledge.db')
CAT = '医保监管'
CAT_DIR = os.path.join(KB_ROOT, 'medical', CAT)
SRC_DIR = os.path.join(KB_ROOT, 'source', '国家医保智能监管知识库规则库')

CN_NUM = '一二三四五六七八九十'


def clean(v):
    if v is None:
        return ''
    s = str(v).replace('\n', ' ').replace('\r', ' ')
    s = re.sub(r'\s+', ' ', s).strip()
    # 去掉 PDF/换行造成的汉字之间多余空格
    s = re.sub(r'(?<=[\u4e00-\u9fff、，。；：（）()]) (?=[\u4e00-\u9fff、，。；：（）()])', '', s)
    return s


def md_esc(s):
    return s.replace('|', '\\|')


def norm_batch(name):
    m = re.search(r'第([%s]+)批' % CN_NUM, name)
    return '第%s批' % m.group(1) if m else ''


def extract_rule(name):
    """从文件名提取规则名"""
    m = re.search(r'[“"](.+?)[”"]', name)
    if m:
        return m.group(1).strip()
    s = re.sub(r'^第[%s]+批' % CN_NUM, '', name)
    s = re.sub(r'^[-—\s]*\d+', '', s)
    s = re.sub(r'规则对应.*$', '', s)
    s = re.sub(r'\.(xlsx|pdf|xls)$', '', s, flags=re.I)
    return s.strip(' .-_　')


def rule_type(rule):
    if '中药饮片' in rule:
        return '中药饮片'
    if '药品' in rule:
        return '药品'
    if '医疗服务项目' in rule or '手术项目' in rule:
        return '医疗服务项目'
    return '其他'


def sheet_table(ws, max_col_limit=64, hdr_pred=None, row_pred=None):
    """读一张 sheet -> (headers, rows)；定位表头行（默认首列为'序号'）"""
    rows = []
    for r in ws.iter_rows(min_col=1, max_col=max_col_limit, values_only=True):
        rows.append(list(r))
    if hdr_pred is None:
        hdr_pred = lambda r: bool(r) and clean(r[0]) == '序号'
    hdr_idx = None
    for i, r in enumerate(rows[:10]):
        if hdr_pred(r):
            hdr_idx = i
            break
    if hdr_idx is None:
        return [], []
    hdr_raw = rows[hdr_idx]
    ncol = 0
    for j, c in enumerate(hdr_raw[:max_col_limit]):
        if clean(c):
            ncol = j + 1
    headers = [clean(c) for c in hdr_raw[:ncol]]
    data = []
    for r in rows[hdr_idx + 1:]:
        r = list(r[:ncol]) + [None] * max(0, ncol - len(r))
        c0 = clean(r[0])
        if not (re.fullmatch(r'\d+', c0) or (row_pred and row_pred(r))):
            continue
        data.append([clean(c) for c in r[:ncol]])
    return headers, data


def main():
    src_root = sys.argv[1]
    dry = '--dry' in sys.argv
    files = sorted(glob.glob(os.path.join(src_root, '**', '*.xlsx'), recursive=True) +
                   glob.glob(os.path.join(src_root, '**', '*.pdf'), recursive=True))
    if not files:
        print('未找到文件'); sys.exit(1)

    os.makedirs(CAT_DIR, exist_ok=True)
    os.makedirs(SRC_DIR, exist_ok=True)
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    if not dry:
        shutil.copytree(src_root, SRC_DIR, dirs_exist_ok=True)

    conn = sqlite3.connect(DB_PATH)
    cur = conn.cursor()
    # 幂等：先清掉本类别旧记录，避免重复入库
    if not dry:
        cur.execute('DELETE FROM insurance_rules WHERE doc_id IN (SELECT id FROM documents WHERE category=?)', (CAT,))
        cur.execute('DELETE FROM insurance_codes WHERE doc_id IN (SELECT id FROM documents WHERE category=?)', (CAT,))
        cur.execute('DELETE FROM documents WHERE category=?', (CAT,))
    cur.execute('''CREATE TABLE IF NOT EXISTS insurance_rules (
        id INTEGER PRIMARY KEY AUTOINCREMENT, doc_id INTEGER, batch TEXT, batch_file TEXT,
        rule_name TEXT, rule_type TEXT, seq INTEGER, item_name TEXT,
        detect_logic TEXT, basis TEXT, code_count INTEGER, extra TEXT)''')
    cur.execute('''CREATE TABLE IF NOT EXISTS insurance_codes (
        id INTEGER PRIMARY KEY AUTOINCREMENT, doc_id INTEGER, batch TEXT, rule_name TEXT,
        seq INTEGER, item_name TEXT, code TEXT)''')
    cur.execute('CREATE INDEX IF NOT EXISTS ix_ir_name ON insurance_rules(item_name)')
    cur.execute('CREATE INDEX IF NOT EXISTS ix_ir_rule ON insurance_rules(rule_name)')
    cur.execute('CREATE INDEX IF NOT EXISTS ix_ic_name ON insurance_codes(item_name)')
    cur.execute('CREATE INDEX IF NOT EXISTS ix_ic_code ON insurance_codes(code)')
    cur.execute('''CREATE TABLE IF NOT EXISTS documents (
        id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL, type TEXT, source TEXT,
        url TEXT, published TEXT, journal TEXT, doi TEXT, authors TEXT, tags TEXT,
        keywords TEXT, category TEXT, file_path TEXT, imported_at TEXT)''')

    now = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    report = []
    tot_kp = tot_code = 0

    for f in files:
        base = os.path.basename(f)
        batch = norm_batch(base)
        rule = extract_rule(base)
        rtype = rule_type(rule)
        title = f'{batch}｜{rule}'
        if base.lower().endswith('.pdf'):
            import pdfplumber
            headers, data, codes = [], [], []
            with pdfplumber.open(f) as pdf:
                for page in pdf.pages:
                    for t in page.extract_tables():
                        for row in t:
                            row = [clean(c) for c in row]
                            if not row or not re.fullmatch(r'\d+', row[0] or ''):
                                continue
                            data.append(row)
            # PDF 行: 序号/名称/检出逻辑/逻辑依据/项目代码/备注
            rows_kp = []
            for row in data:
                row = row + [''] * (6 - len(row))
                name, logic, basis, code, note = row[1], row[2], row[3], row[4], row[5]
                rows_kp.append((int(row[0]), name, logic, basis, '', {'项目代码': code, '备注': note}))
                if code:
                    codes.append((int(row[0]), name, code))
            headers = ['序号', '医疗服务项目名称', '检出逻辑', '逻辑依据（同切口组别名称）', '项目代码供参考', '备注']
            kp_rows = rows_kp
        else:
            wb = load_workbook(f, read_only=True, data_only=True)
            ws1 = wb.worksheets[0]
            headers, data = sheet_table(ws1)
            kp_rows = []
            for r in data:
                seq = int(r[0])
                name = r[1] if len(r) > 1 else ''
                logic = next((r[j] for j, h in enumerate(headers) if '检出逻辑' in h), '')
                basis = next((r[j] for j, h in enumerate(headers) if '逻辑依据' in h), '')
                ccount = next((r[j] for j, h in enumerate(headers) if '代码数量' in h), '')
                extra = {headers[j]: r[j] for j in range(len(headers)) if r[j]
                         and j > 1 and '检出逻辑' not in headers[j] and '逻辑依据' not in headers[j]
                         and '代码数量' not in headers[j] and headers[j]}
                cc = int(ccount) if re.fullmatch(r'\d+', ccount or '') else None
                kp_rows.append((seq, name, logic, basis, cc, extra))
            codes = []
            if len(wb.worksheets) > 1:
                _, cdata = sheet_table(
                    wb.worksheets[1],
                    hdr_pred=lambda r: any(clean(c) == '序号' for c in r)
                    and any('代码' in clean(c) for c in r),
                    row_pred=lambda r: len(r) > 3 and clean(r[3]) != '')
                last_seq, last_name = None, ''
                for r in cdata:
                    if len(r) < 4:
                        continue
                    if r[0] and re.fullmatch(r'\d+', r[0]):
                        last_seq = int(r[0])
                        last_name = r[1] or last_name
                    if r[3]:
                        codes.append((last_seq, r[1] or last_name, r[3]))
            wb.close()

        # ---- markdown ----
        safe_rule = re.sub(r'[/\\:*?"<>|]', '', rule)
        md_name = f'{batch}-{safe_rule}.md'
        md_path = os.path.join(CAT_DIR, md_name)
        tags = json.dumps([CAT, rtype, batch], ensure_ascii=False)
        hdr_line = '| ' + ' | '.join(md_esc(h) for h in headers) + ' |'
        sep_line = '|' + '|'.join(['---'] * len(headers)) + '|'
        body_rows = []
        for (seq, name, logic, basis, cc, extra) in kp_rows:
            cells = []
            for j, h in enumerate(headers):
                if j == 0:
                    cells.append(str(seq))
                elif j == 1:
                    cells.append(md_esc(name))
                elif '检出逻辑' in h:
                    cells.append(md_esc(logic))
                elif '逻辑依据' in h:
                    cells.append(md_esc(basis))
                elif '代码数量' in h:
                    cells.append(str(cc) if cc is not None else '')
                else:
                    cells.append(md_esc(str(extra.get(h, ''))))
            body_rows.append('| ' + ' | '.join(cells) + ' |')

        md = f"""---
title: "{batch} {rule}（国家医保智能监管规则知识点）"
type: 规则库
category: {CAT}
source: "国家医疗保障局 医保智能监管知识库、规则库"
url: ""
published: ""
authors: "国家医疗保障局"
tags: {tags}
keywords: "{rule},{rtype},医保智能监管,规则知识点,{batch}"
imported_at: {now}
---

# {batch} {rule}（国家医保智能监管规则知识点）

> **来源**: 国家医疗保障局《国家医保智能监管知识库、规则库》
> **批次**: {batch}　**规则类型**: {rtype}　**规则名称**: {rule}
> **知识点数量**: {len(kp_rows)}　**对应{('项目/药品代码' if codes else '代码')}数量**: {len(codes)}
> **原始文件**: `source/国家医保智能监管知识库规则库/{base}`
> **机器可检索副本**: SQLite `knowledge.db` 表 `insurance_rules`（知识点）{('/ `insurance_codes`（代码明细）' if codes else '')}

## 规则说明

| 项目 | 内容 |
|------|------|
| 规则名称 | {rule} |
| 批次 | {batch} |
| 规则类型 | {rtype} |
| 知识点条数 | {len(kp_rows)} |
| 对应代码条数 | {len(codes)} |

## 知识点明细（{len(kp_rows)} 条）

{hdr_line}
{sep_line}
{chr(10).join(body_rows)}
"""
        if not dry:
            with open(md_path, 'w') as fh:
                fh.write(md)
            cur.execute('''INSERT INTO documents
                (title, type, source, url, published, journal, doi, authors, tags, keywords,
                 category, file_path, imported_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)''',
                (f'{batch} {rule}（国家医保智能监管规则知识点）', '规则库',
                 '国家医疗保障局 医保智能监管知识库、规则库', '', '', '', '', '国家医疗保障局',
                 ','.join([CAT, rtype, batch]), f'{rule},{rtype},医保智能监管,规则知识点,{batch}',
                 CAT, md_path, now))
            doc_id = cur.lastrowid
            cur.executemany('''INSERT INTO insurance_rules
                (doc_id,batch,batch_file,rule_name,rule_type,seq,item_name,detect_logic,basis,code_count,extra)
                VALUES (?,?,?,?,?,?,?,?,?,?,?)''',
                [(doc_id, batch, base, rule, rtype, s, n, l, b, c,
                  json.dumps(e, ensure_ascii=False)) for (s, n, l, b, c, e) in kp_rows])
            if codes:
                cur.executemany('''INSERT INTO insurance_codes
                    (doc_id,batch,rule_name,seq,item_name,code) VALUES (?,?,?,?,?,?)''',
                    [(doc_id, batch, rule, s, n, c) for (s, n, c) in codes])
        tot_kp += len(kp_rows)
        tot_code += len(codes)
        report.append((batch, rule, rtype, len(kp_rows), len(codes), md_name))

    if not dry:
        conn.commit()
        n_doc = cur.execute('SELECT COUNT(*) FROM documents').fetchone()[0]
        n_r = cur.execute('SELECT COUNT(*) FROM insurance_rules').fetchone()[0]
        n_c = cur.execute('SELECT COUNT(*) FROM insurance_codes').fetchone()[0]
    else:
        n_doc = n_r = n_c = '(dry)'

    print('批次  规则  类型  知识点  代码  文件')
    for b, r, t, k, c, m in report:
        print(f'  {b:<6} {r:<24} {t:<7} {k:>5} {c:>6}  {m}')
    print(f'\n本次: 批次文件 {len(report)} 个, 知识点 {tot_kp} 条, 代码 {tot_code} 条')
    print(f'库内累计: documents={n_doc}, insurance_rules={n_r}, insurance_codes={n_c}')
    conn.close()


if __name__ == '__main__':
    main()
