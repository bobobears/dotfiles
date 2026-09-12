#!/usr/bin/env python3
"""安徽省医疗服务价格项目 皖医保发〔2025〕18号/22号 入库脚本（幂等）。
下载附件 → .doc转.docx(LibreOffice) → python-docx解析表格 → SQLite yb_price_items + markdown。
用法: python3 kb_import_ahprice.py [--dry]
"""
import os, re, sys, glob, shutil, sqlite3, subprocess, datetime
import docx

BASE = "https://ybj.ah.gov.cn"
UA = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120.0.0.0 Safari/537.36"
KB = os.path.expanduser("~/private_db/knowledge")
DB = os.path.join(KB, "db", "knowledge.db")
MD_DIR = os.path.join(KB, "medical", "医疗服务价格")
SRC_DIR = os.path.join(KB, "source", "安徽省医疗服务价格项目2025-18号22号")
WORK = "/tmp/ah_yz/work"

DOCS = [
    {
        "docno": "皖医保发〔2025〕18号",
        "title": "关于规范整合综合诊查、超声检查、精神治疗、放射治疗及康复类医疗服务价格项目的通知",
        "url": f"{BASE}/public/7071/150361601.html",
        "published": "2025-12-15",
        "category": "综合诊查/超声检查/精神治疗/放射治疗/康复类",
        "atts": [  # (文件名, 相对路径)
            ("整合后综合诊查类医疗服务价格项目表.doc", "/group6/M00/0E/AE/wKg8Bmk_f_-AIgtQAARKAG7ySUI294.doc"),
            ("整合后超声检查类医疗服务价格项目表.doc", "/group6/M00/0E/AE/wKg8Bmk_gA2APR0KAANqAJCq3bA714.doc"),
            ("整合后精神治疗类医疗服务价格项目表.doc", "/group6/M00/0E/AE/wKg8Bmk_gCKABkiXAAFCAIs9zk0835.doc"),
            ("整合后放射治疗类医疗服务价格项目表.doc", "/group6/M00/0E/AE/wKg8Bmk_gC-AXm3MAAK0AOigfIM420.doc"),
            ("整合后康复类医疗服务价格项目表.doc", "/group6/M00/0E/AE/wKg8Bmk_gDqAKUsuAAOAAJ17P2o259.doc"),
            ("废止医疗服务价格项目表(18号).doc", "/group6/M00/0E/AE/wKg8Bmk_gMuAOVgHAA7aAFELaVY531.doc"),
            ("整合后医疗服务价格项目映射关系表(18号).doc", "/group6/M00/0E/AE/wKg8Bmk_gNuAA2J-AAdKACgtq3Y262.doc"),
        ],
    },
    {
        "docno": "皖医保发〔2025〕22号",
        "title": "关于规范整合呼吸系统、神经系统等十二类医疗服务价格项目的通知",
        "url": f"{BASE}/public/7071/150415601.html",
        "published": "2025-12-29",
        "category": "血液/神经/耳鼻喉/呼吸/心血管/泌尿/妇科/骨骼肌肉/疝甲乳/体被/物理治疗/美容整形类",
        "atts": [
            ("整合后血液系统类医疗服务价格项目表.docx", "/group6/M00/0E/E8/wKg8Bml8qyqAJZPnAAUUHAtXwro14.docx"),
            ("整合后神经系统类医疗服务价格项目表.docx", "/group6/M00/0E/E8/wKg8Bml8qz2AQ0qMAAcQ2aiCEd085.docx"),
            ("整合后耳鼻喉科类医疗服务价格项目表.docx", "/group6/M00/0E/E8/wKg8Bml8q1mAUEtdAAgh51-ZEY865.docx"),
            ("整合后呼吸系统类医疗服务价格项目表.docx", "/group6/M00/0E/E8/wKg8Bml8q2eAWRgWAAZ34tTUFxc43.docx"),
            ("整合后心血管系统类医疗服务价格项目表.docx", "/group6/M00/0E/E8/wKg8Bml8q3SADmXrAAf-RvGQrTw54.docx"),
            ("整合后泌尿系统类医疗服务价格项目表.docx", "/group6/M00/0E/E8/wKg8Bml8q36AA6RZAAa15EsmaNc18.docx"),
            ("整合后妇科类医疗服务价格项目表.docx", "/group6/M00/0E/E8/wKg8Bml8q4yAcLUDAAYs5BMoD1I16.docx"),
            ("整合后骨骼肌肉系统类医疗服务价格项目表.docx", "/group6/M00/0E/E8/wKg8Bml8q56AB52MAAd0NTmDycQ32.docx"),
            ("整合后疝甲乳类医疗服务价格项目表.docx", "/group6/M00/0E/E8/wKg8Bml8q6uAVtyDAAWsScGe1jw29.docx"),
            ("整合后体被系统类医疗服务价格项目表.docx", "/group6/M00/0E/E8/wKg8Bml8q7mAS1yQAAYkVyzOK5E32.docx"),
            ("整合后物理治疗类医疗服务价格项目表.docx", "/group6/M00/0E/E8/wKg8Bml8q8SAFWMnAAVeCIHhs5U81.docx"),
            ("整合后美容整形类医疗服务价格项目表.docx", "/group6/M00/0E/E8/wKg8Bml8q8-Abl32AAcYWGpIvs071.docx"),
            ("废止医疗服务价格项目表(22号).docx", "/group6/M00/0E/E8/wKg8Bml8q9uAcegwAA5dhNMygCY47.docx"),
        ],
    },
]

def clean(s):
    if s is None: return ""
    s = str(s).replace("\xa0", " ").strip()
    s = re.sub(r"\s+", " ", s)
    return s

def download(url, path, retries=3):
    for i in range(retries):
        r = subprocess.run(["curl", "-sL", "--connect-timeout", "25", "--max-time", "180",
                            "-A", UA, url, "-o", path], capture_output=True)
        if os.path.exists(path) and os.path.getsize(path) > 1000:
            return True
        print(f"      (下载重试 {i+1}/{retries})")
    return False

def doc_to_docx(docpath):
    outdir = os.path.dirname(docpath)
    r = subprocess.run(["soffice", "--headless", "--convert-to", "docx", docpath,
                        "--outdir", outdir], capture_output=True, timeout=180)
    target = re.sub(r"\.doc$", ".docx", docpath)
    return os.path.exists(target)

def parse_table(docx_path):
    """返回 (headers, rows)。rows 为 list[list[str]]。"""
    d = docx.Document(docx_path)
    if not d.tables:
        return None, []
    t = d.tables[0]
    headers = [clean(c.text) for c in t.rows[0].cells]
    rows = []
    for r in t.rows[1:]:
        cells = [clean(c.text) for c in r.cells]
        # 去尾部空列
        while cells and cells[-1] == "":
            cells.pop()
        if any(cells):
            rows.append(cells)
    return headers, rows

def main():
    dry = "--dry" in sys.argv
    os.makedirs(WORK, exist_ok=True)
    os.makedirs(MD_DIR, exist_ok=True)
    os.makedirs(SRC_DIR, exist_ok=True)

    all_items = []   # 结构化记录
    doc_meta = {}    # docno -> {title,url,published,category,n_items}

    for doc in DOCS:
        print(f"\n### {doc['docno']}")
        n_doc = 0
        for fname, rel in doc["atts"]:
            ext = os.path.splitext(fname)[1]
            local_raw = os.path.join(WORK, f"{doc['docno'][:6]}_{fname}")
            url = BASE + rel
            if not (os.path.exists(local_raw) and os.path.getsize(local_raw) > 1000):
                ok = download(url, local_raw)
                print(f"  {'OK ' if ok else 'FAIL'} {fname} ({os.path.getsize(local_raw) if os.path.exists(local_raw) else 0}B)")
            if not (os.path.exists(local_raw) and os.path.getsize(local_raw) > 1000):
                print(f"      !! 跳过（下载失败）")
                continue
            # 转 docx
            work_docx = re.sub(r"\.doc$", ".docx", local_raw)
            if ext == ".doc":
                if not (os.path.exists(work_docx) and os.path.getsize(work_docx) > 500):
                    ok = doc_to_docx(local_raw)
                    print(f"      -> {work_docx.split('/')[-1]} {'OK' if ok else 'FAIL'}")
            # source 备份（保留原始格式）
            shutil.copy2(local_raw, os.path.join(SRC_DIR, fname))

            if not (os.path.exists(work_docx) and os.path.getsize(work_docx) > 500):
                print(f"      !! 跳过（转换失败）")
                continue
            headers, rows = parse_table(work_docx)
            is_abolish = "废止" in fname
            is_map = "映射" in fname
            print(f"      {fname}: {len(rows)}行 {'[废止表]' if is_abolish else '[映射表-不入库]' if is_map else ''}")

            # 映射关系表是6列对照(无单价)，不进价格库，仅 source 备份 + md 说明
            if is_map:
                continue

            for row in rows:
                rec = {"docno": doc["docno"], "file": os.path.splitext(fname)[0],
                       "category": doc["category"]}
                # 按列名对齐（表头单元格可能含内部空格，如"价格 （元）"；废止表用"项目编码/项目内涵"）
                if headers and len(row) >= 3:
                    h = [re.sub(r"\s+", "", x) for x in headers]
                    def g(*names):
                        for nm in names:
                            key = re.sub(r"\s+", "", nm)
                            for i, hh in enumerate(h):
                                if hh == key and i < len(row):
                                    return row[i]
                        return ""
                    rec["seq"] = g("序号")
                    rec["code"] = g("项目代码", "项目编码")
                    rec["name"] = g("项目名称")
                    rec["output"] = g("服务产出", "项目内涵")
                    rec["price_comp"] = g("价格构成")
                    rec["unit"] = g("计价单位")
                    rec["price"] = g("价格（元）")
                    rec["note"] = g("计价说明") + (("；除外：" + g("除外内容")) if g("除外内容") else "")
                    rec["pay_class"] = g("支付分类")
                    rec["stat_cat"] = g("统计/分类", "分类")
                else:
                    # 废止表等：尽量按位置
                    rec.update({"seq": row[0] if len(row)>0 else "", "code": row[1] if len(row)>1 else "",
                                "name": row[2] if len(row)>2 else ""})
                rec["is_abolish"] = is_abolish
                all_items.append(rec)
            n_doc += len(rows)

        doc_meta[doc["docno"]] = {"title": doc["title"], "url": doc["url"],
                                  "published": doc["published"], "category": doc["category"]}

    print(f"\n=== 合计解析 {len(all_items)} 条记录 ===")
    if dry:
        # 抽样打印
        for r in all_items[:3]:
            print("  ", r.get("code"), r.get("name"), r.get("price"), r.get("unit"))
        print("[dry] 未写库")
        return

    # ---- 建表 + 幂等清理 ----
    conn = sqlite3.connect(DB)
    cur = conn.cursor()
    cur.execute("""CREATE TABLE IF NOT EXISTS yb_price_items (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        docno TEXT, file TEXT, category TEXT, seq TEXT, code TEXT, name TEXT,
        output TEXT, price_comp TEXT, unit TEXT, price TEXT, note TEXT,
        pay_class TEXT, stat_cat TEXT, is_abolish INTEGER DEFAULT 0)""")
    cur.execute("DELETE FROM yb_price_items WHERE docno IN ('皖医保发〔2025〕18号','皖医保发〔2025〕22号')")

    for r in all_items:
        cur.execute("""INSERT INTO yb_price_items (docno,file,category,seq,code,name,output,price_comp,unit,price,note,pay_class,stat_cat,is_abolish)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (r["docno"], r["file"], r["category"], r.get("seq",""), r.get("code",""), r.get("name",""),
             r.get("output",""), r.get("price_comp",""), r.get("unit",""), r.get("price",""),
             r.get("note",""), r.get("pay_class",""), r.get("stat_cat",""), 1 if r["is_abolish"] else 0))

    # ---- documents 登记（幂等）----
    now = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    for docno, m in doc_meta.items():
        cur.execute("DELETE FROM documents WHERE title=? AND category='医疗服务价格'", (docno + " " + m["title"],))
        n_items = sum(1 for r in all_items if r["docno"] == docno and not r["is_abolish"])
        md_path = os.path.join(MD_DIR, f"{docno}.md")
        cur.execute("""INSERT INTO documents (title,type,source,url,published,tags,category,file_path,imported_at)
            VALUES (?,?,?,?,?,?,?,?,?)""",
            (f"{docno} {m['title']}", "行政规范性文件", "安徽省医疗保障局", m["url"], m["published"],
             "医疗服务价格;项目整合;" + docno, "医疗服务价格", md_path, now))

    conn.commit()
    print(f"yb_price_items 写入 {len(all_items)} 条")
    print("documents 登记 2 条")

    # ---- markdown ----
    for docno, m in doc_meta.items():
        items = [r for r in all_items if r["docno"] == docno]
        live = [r for r in items if not r["is_abolish"]]
        abol = [r for r in items if r["is_abolish"]]
        lines = [f"# {docno} {m['title']}", "",
                 f"- **发文机关**：安徽省医疗保障局",
                 f"- **发布日期**：{m['published']}",
                 f"- **原文链接**：{m['url']}",
                 f"- **覆盖类别**：{m['category']}",
                 f"- **整合后项目数**：{len(live)}；废止项目数：{len(abol)}", "",
                 "> 数据来源：安徽省医保局官网部门文件栏目附件（.doc/.docx 价格项目表）。",
                 "> 单价为省级基准价，具体执行以各市/机构公示为准。", ""]
        # 按 file 分组
        files = {}
        for r in items:
            files.setdefault(r["file"], []).append(r)
        for fname, rows in files.items():
            lines.append(f"## {fname}")
            lines.append("")
            if any("废止" in fname for _ in [0]):
                pass
            hdr = ["序号","项目代码","项目名称","计价单位","价格(元)","计价说明"]
            lines.append("| " + " | ".join(hdr) + " |")
            lines.append("|" + "---|"*len(hdr))
            for r in rows:
                cells = [r.get("seq",""), r.get("code",""), r.get("name",""),
                         r.get("unit",""), r.get("price",""), (r.get("note","") or "").replace("\n"," ")[:120]]
                lines.append("| " + " | ".join(c.replace("|","/") for c in cells) + " |")
            lines.append("")
        with open(os.path.join(MD_DIR, f"{docno}.md"), "w", encoding="utf-8") as f:
            f.write("\n".join(lines))
    print(f"markdown 写入 {MD_DIR}")

if __name__ == "__main__":
    main()
