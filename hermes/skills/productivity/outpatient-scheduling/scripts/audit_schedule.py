#!/usr/bin/env python3
"""排班表审计（通用，任意月份）

用法: python3 audit_schedule.py <排班表.xlsx>

假设表结构：行1=标题，行3=表头(日期|星期|早班|晚班|备班|休假)，行4起为数据；
周末/周分隔行为合并单元格，第2列（星期）为空，自动跳过。

校验项：
  1. 同日早班≠晚班
  2. 无连续两天相同班次（早/晚分开查）
  3. 休假顺序：周六休=周五早班人，周日休=周五晚班人
  4. 周一~周四有备班；周五~周日无备班
  5. 每日早/晚班均不为空
  6. 输出每人班次统计（用于确认均衡）

退出码：0=全部通过，1=有问题（可接 CI/脚本判断）
"""
import sys
from collections import Counter

from openpyxl import load_workbook

WEEKDAYS = ('一', '二', '三', '四', '五', '六', '日')


def load_rows(path):
    ws = load_workbook(path).active
    rows = []
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        date, week = r[0], r[1]
        if not date:
            continue
        wk = str(week).strip() if week else ''
        if wk not in WEEKDAYS:          # 标题行/周分隔行
            continue
        rows.append({
            'date': str(date).strip(),
            'week': wk,
            'early': (str(r[2]).strip() if len(r) > 2 and r[2] else ''),
            'late': (str(r[3]).strip() if len(r) > 3 and r[3] else ''),
            'backup': (str(r[4]).strip() if len(r) > 4 and r[4] else ''),
            'off': (str(r[5]).strip() if len(r) > 5 and r[5] else ''),
        })
    return rows


def audit(rows):
    errs = []

    # 1. 同日早≠晚
    for x in rows:
        if x['early'] and x['early'] == x['late']:
            errs.append(f"{x['date']} 早班=晚班={x['early']}")

    # 2. 连续两天相同班次（逐日相邻比较）
    for prev, cur in zip(rows, rows[1:]):
        if cur['early'] and cur['early'] == prev['early']:
            errs.append(f"{prev['date']}→{cur['date']} 连续早班: {cur['early']}")
        if cur['late'] and cur['late'] == prev['late']:
            errs.append(f"{prev['date']}→{cur['date']} 连续晚班: {cur['late']}")

    # 3. 休假顺序：切分自然周（遇到「周一」就新开一组）
    groups, cur = [], []
    for x in rows:
        if x['week'] == '一' and cur:
            groups.append(cur)
            cur = [x]
        else:
            cur.append(x)
    if cur:
        groups.append(cur)

    for g in groups:
        fri = next((d for d in g if d['week'] == '五'), None)
        sat = next((d for d in g if d['week'] == '六'), None)
        sun = next((d for d in g if d['week'] == '日'), None)
        if fri and sat and sat['off'] and fri['early'] and sat['off'] != fri['early']:
            errs.append(f"{sat['date']} 周六休假应={fri['early']}，实际={sat['off']}")
        if fri and sun and sun['off'] and fri['late'] and sun['off'] != fri['late']:
            errs.append(f"{sun['date']} 周日休假应={fri['late']}，实际={sun['off']}")

    # 4. 备班覆盖
    for x in rows:
        has = bool(x['backup']) and x['backup'] != '-'
        if x['week'] in ('一', '二', '三', '四') and not has:
            errs.append(f"{x['date']} 周一~周四缺备班")
        if x['week'] in ('五', '六', '日') and has:
            errs.append(f"{x['date']} 周五~周日不应有备班，实际={x['backup']}")

    # 5. 早晚班覆盖
    for x in rows:
        if not x['early'] or not x['late']:
            errs.append(f"{x['date']} 早/晚班缺失")

    return errs


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    path = sys.argv[1]
    rows = load_rows(path)
    print(f"解析 {len(rows)} 天：{rows[0]['date']} ~ {rows[-1]['date']}\n")

    errs = audit(rows)
    print("=== 规则审计 ===")
    if errs:
        for e in errs:
            print("  ❌", e)
    else:
        print("  ✅ 全部通过")

    stat = {}
    for x in rows:
        for role in ('early', 'late', 'backup'):
            v = x[role]
            if v and v != '-':
                stat.setdefault(v, Counter())[role] += 1
    print("\n=== 每人班次统计 ===")
    for name, c in stat.items():
        print(f"  {name}: 早班{c['early']} 晚班{c['late']} 备班{c['backup']} "
              f"合计{sum(c.values())}")

    sys.exit(1 if errs else 0)


if __name__ == '__main__':
    main()
