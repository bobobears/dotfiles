"""
2026年9月排班表生成脚本
基于门诊排班skill规则，从8月末状态衔接推算

8月末衔接状态 (来自 gen_schedule_aug.py):
  8/31(一): E=孙闪闪, L=二院派遣, 备=陈东升
  自然周 8/31-9/6 跨月，备班保持陈东升不变
  备班轮换顺序: 二院派遣(姜湛乾) → 陈东升 → 孙闪闪

推导过程:
──────── Week A (8/31-9/6): backup=陈东升, pair={二院派遣, 孙闪闪} ────────
  8/31(一): E=孙闪闪, L=二院派遣 (八月末状态)
  9/1(二): E=二院派遣, L=孙闪闪
  9/2(三): E=孙闪闪, L=二院派遣
  9/3(四): E=二院派遣, L=孙闪闪
  9/4(五): E=孙闪闪, L=二院派遣, 休=陈东升 (备班人周五休)
  9/5(六): 周五晚(二院派遣)→早, 周五休(陈东升)→晚 → E二院派遣, L陈东升, 休孙闪闪
  9/6(日): 周六晚(陈东升)→早, 周六休(孙闪闪)→晚 → E陈东升, L孙闪闪, 休二院派遣

──────── Week B (9/7-9/13): backup=孙闪闪 ────────
  Sun(9/6)晚=孙闪闪 IS 本周备班 → 周日休班人(二院派遣)做周一早
  9/7(一): E=二院派遣, L=陈东升, 备=孙闪闪
  9/8(二): E=陈东升, L=二院派遣, 备=孙闪闪
  9/9(三): E=二院派遣, L=陈东升, 备=孙闪闪
  9/10(四): E=陈东升, L=二院派遣, 备=孙闪闪
  9/11(五): E=二院派遣, L=陈东升, 休=孙闪闪
  9/12(六): 周五晚(陈东升)→早, 周五休(孙闪闪)→晚 → E陈东升, L孙闪闪, 休二院派遣
  9/13(日): 周六晚(孙闪闪)→早, 周六休(二院派遣)→晚 → E孙闪闪, L二院派遣, 休陈东升

──────── Week C (9/14-9/20): backup=二院派遣 ────────
  Sun(9/13)晚=二院派遣 IS 本周备班 → 周日休班人(陈东升)做周一早
  9/14(一): E=陈东升, L=孙闪闪, 备=二院派遣
  9/15(二): E=孙闪闪, L=陈东升, 备=二院派遣
  9/16(三): E=陈东升, L=孙闪闪, 备=二院派遣
  9/17(四): E=孙闪闪, L=陈东升, 备=二院派遣
  9/18(五): E=陈东升, L=孙闪闪, 休=二院派遣
  9/19(六): 周五晚(孙闪闪)→早, 周五休(二院派遣)→晚 → E孙闪闪, L二院派遣, 休陈东升
  9/20(日): 周六晚(二院派遣)→早, 周六休(陈东升)→晚 → E二院派遣, L陈东升, 休孙闪闪

──────── Week D (9/21-9/27): backup=陈东升 ────────
  Sun(9/20)晚=陈东升 IS 本周备班 → 周日休班人(孙闪闪)做周一早
  9/21(一): E=孙闪闪, L=二院派遣, 备=陈东升
  9/22(二): E=二院派遣, L=孙闪闪, 备=陈东升
  9/23(三): E=孙闪闪, L=二院派遣, 备=陈东升
  9/24(四): E=二院派遣, L=孙闪闪, 备=陈东升
  9/25(五): E=孙闪闪, L=二院派遣, 休=陈东升
  9/26(六): 周五晚(二院派遣)→早, 周五休(陈东升)→晚 → E二院派遣, L陈东升, 休孙闪闪
  9/27(日): 周六晚(陈东升)→早, 周六休(孙闪闪)→晚 → E陈东升, L孙闪闪, 休二院派遣

──────── Week E (9/28-9/30): backup=孙闪闪 ────────
  Sun(9/27)晚=孙闪闪 IS 本周备班 → 周日休班人(二院派遣)做周一早
  9/28(一): E=二院派遣, L=陈东升, 备=孙闪闪
  9/29(二): E=陈东升, L=二院派遣, 备=孙闪闪
  9/30(三): E=二院派遣, L=陈东升, 备=孙闪闪
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ============================================================
# Build schedule data
# (date, weekday, early, late, backup, off, is_weekend)
# ============================================================
sched = [
    ("W1", "第1周 (8/31-9/6)  |  备班人员：陈东升（跨月延续八月末班次）"),

    # 8/31(一): 八月末衔接状态，作为上下文行
    ("8/31", "一", "孙闪闪", "二院派遣", "陈东升", "", False),
    ("9/1", "二", "二院派遣", "孙闪闪", "陈东升", "", False),
    ("9/2", "三", "孙闪闪", "二院派遣", "陈东升", "", False),
    ("9/3", "四", "二院派遣", "孙闪闪", "陈东升", "", False),
    # 9/4(五): 备班人陈东升休假
    ("9/4", "五", "孙闪闪", "二院派遣", "", "陈东升", True),
    # 9/5(六): 周五晚(二院派遣)→早, 周五休(陈东升)→晚
    ("9/5", "六", "二院派遣", "陈东升", "", "孙闪闪", True),
    # 9/6(日): 周六晚(陈东升)→早, 周六休(孙闪闪)→晚
    ("9/6", "日", "陈东升", "孙闪闪", "", "二院派遣", True),

    ("W2", "第2周 (9/7-9/13)  |  备班人员：孙闪闪"),

    # 9/7(一): Sun晚=孙闪闪IS备班 → 周日休班人(二院派遣)做周一早
    ("9/7", "一", "二院派遣", "陈东升", "孙闪闪", "", False),
    ("9/8", "二", "陈东升", "二院派遣", "孙闪闪", "", False),
    ("9/9", "三", "二院派遣", "陈东升", "孙闪闪", "", False),
    ("9/10", "四", "陈东升", "二院派遣", "孙闪闪", "", False),
    # 9/11(五): 备班人孙闪闪休假
    ("9/11", "五", "二院派遣", "陈东升", "", "孙闪闪", True),
    # 9/12(六): 周五晚(陈东升)→早, 周五休(孙闪闪)→晚
    ("9/12", "六", "陈东升", "孙闪闪", "", "二院派遣", True),
    # 9/13(日): 周六晚(孙闪闪)→早, 周六休(二院派遣)→晚
    ("9/13", "日", "孙闪闪", "二院派遣", "", "陈东升", True),

    ("W3", "第3周 (9/14-9/20)  |  备班人员：二院派遣"),

    # 9/14(一): Sun晚=二院派遣IS备班 → 周日休班人(陈东升)做周一早
    ("9/14", "一", "陈东升", "孙闪闪", "二院派遣", "", False),
    ("9/15", "二", "孙闪闪", "陈东升", "二院派遣", "", False),
    ("9/16", "三", "陈东升", "孙闪闪", "二院派遣", "", False),
    ("9/17", "四", "孙闪闪", "陈东升", "二院派遣", "", False),
    # 9/18(五): 备班人二院派遣休假
    ("9/18", "五", "陈东升", "孙闪闪", "", "二院派遣", True),
    # 9/19(六): 周五晚(孙闪闪)→早, 周五休(二院派遣)→晚
    ("9/19", "六", "孙闪闪", "二院派遣", "", "陈东升", True),
    # 9/20(日): 周六晚(二院派遣)→早, 周六休(陈东升)→晚
    ("9/20", "日", "二院派遣", "陈东升", "", "孙闪闪", True),

    ("W4", "第4周 (9/21-9/27)  |  备班人员：陈东升"),

    # 9/21(一): Sun晚=陈东升IS备班 → 周日休班人(孙闪闪)做周一早
    ("9/21", "一", "孙闪闪", "二院派遣", "陈东升", "", False),
    ("9/22", "二", "二院派遣", "孙闪闪", "陈东升", "", False),
    ("9/23", "三", "孙闪闪", "二院派遣", "陈东升", "", False),
    ("9/24", "四", "二院派遣", "孙闪闪", "陈东升", "", False),
    # 9/25(五): 备班人陈东升休假
    ("9/25", "五", "孙闪闪", "二院派遣", "", "陈东升", True),
    # 9/26(六): 周五晚(二院派遣)→早, 周五休(陈东升)→晚
    ("9/26", "六", "二院派遣", "陈东升", "", "孙闪闪", True),
    # 9/27(日): 周六晚(陈东升)→早, 周六休(孙闪闪)→晚
    ("9/27", "日", "陈东升", "孙闪闪", "", "二院派遣", True),

    ("W5", "第5周 (9/28-9/30)  |  备班人员：孙闪闪（仅周一至周三）"),

    # 9/28(一): Sun晚=孙闪闪IS备班 → 周日休班人(二院派遣)做周一早
    ("9/28", "一", "二院派遣", "陈东升", "孙闪闪", "", False),
    ("9/29", "二", "陈东升", "二院派遣", "孙闪闪", "", False),
    ("9/30", "三", "二院派遣", "陈东升", "孙闪闪", "", False),
]

# ============================================================
# Styles (与7月/8月脚本一致)
# ============================================================
wb = Workbook()
ws = wb.active
ws.title = "排班表"

hf = Font(name="Arial", bold=True, size=11, color="FFFFFF")
hfl = PatternFill("solid", fgColor="2C2C2A")
wlf = PatternFill("solid", fgColor="D3D1C7")
dfl = PatternFill("solid", fgColor="F1EFE8")
wf = PatternFill("solid", fgColor="FAEEDA")
wdf = PatternFill("solid", fgColor="FAC775")
ef = Font(name="Arial", bold=True, size=11, color="0C447C")
lf = Font(name="Arial", bold=True, size=11, color="712B13")
bf = Font(name="Arial", bold=True, size=11, color="085041")
offf = Font(name="Arial", bold=True, size=11, color="A32D2D")
df = Font(name="Arial", bold=True, size=11, color="888780")
lab = Font(name="Arial", bold=True, size=11, color="2C2C2A")
nf = Font(name="Arial", size=11, color="444441")
th = Border(
    left=Side(style="thin", color="B4B2A9"),
    right=Side(style="thin", color="B4B2A9"),
    top=Side(style="thin", color="B4B2A9"),
    bottom=Side(style="thin", color="B4B2A9"),
)
ca = Alignment(horizontal="center", vertical="center")
la = Alignment(horizontal="left", vertical="center")

# === Title ===
ws.merge_cells("A1:F1")
ws["A1"] = "2026年9月工作排班表（9月1日 — 9月30日）"
ws["A1"].font = Font(name="Arial", bold=True, size=14, color="2C2C2A")
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 32

# === Headers ===
headers = ["日期", "星期", "早班\n08:00-14:00", "晚班\n14:00-20:00", "备班\n08:00-12:00\n+14:00-17:30", "休假"]
ws.row_dimensions[3].height = 40
for i, h in enumerate(headers, 1):
    c = ws.cell(row=3, column=i, value=h)
    c.font = hf
    c.fill = hfl
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = th

# === Write schedule ===
row = 4
for item in sched:
    if isinstance(item[0], str) and item[0].startswith("W"):
        # Week header row
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        c = ws.cell(row=row, column=1, value=item[1])
        c.font = lab
        c.fill = wlf
        c.alignment = la
        c.border = th
        for cc in range(2, 7):
            ws.cell(row=row, column=cc).fill = wlf
            ws.cell(row=row, column=cc).border = th
        ws.row_dimensions[row].height = 24
        row += 1
    else:
        date, wd, early, late, backup, off_p, is_we = item
        data = [date, wd, early, late, backup, off_p if off_p else "-"]
        for c_idx, val in enumerate(data, 1):
            cell = ws.cell(row=row, column=c_idx, value=val)
            cell.border = th
            cell.alignment = ca
            if c_idx == 1:
                cell.fill = wdf if is_we else dfl
                cell.font = lab
            elif c_idx == 2:
                cell.fill = wf if is_we else dfl
                cell.font = lab
            else:
                if is_we:
                    cell.fill = wf
                # 派遣人员标红底黑字
                rfl = PatternFill("solid", fgColor="FF6666")
                if val == "二院派遣":
                    cell.font = nf
                    cell.fill = rfl
                elif c_idx == 3:
                    cell.font = ef  # 早班=蓝色
                elif c_idx == 4:
                    cell.font = lf  # 晚班=橙色
                elif c_idx == 5:
                    cell.font = df if val == "-" else bf  # 备班=绿色
                elif c_idx == 6:
                    cell.font = df if val == "-" else offf  # 休假=红色
        ws.row_dimensions[row].height = 24
        row += 1

# === Column widths ===
ws.column_dimensions["A"].width = 10
ws.column_dimensions["B"].width = 8
ws.column_dimensions["C"].width = 18
ws.column_dimensions["D"].width = 18
ws.column_dimensions["E"].width = 24
ws.column_dimensions["F"].width = 42

fp = "/home/bobobears/下载/排班表_2026年9月.xlsx"
wb.save(fp)
print(f"✅ 已生成：{fp}")

# ============================================================
# Audit: triple verification (含8/31衔接行，覆盖跨月边界)
# ============================================================
print("\n═══════════ 审计验证 ═══════════")

errors = []
all_items = [it for it in sched if not (isinstance(it[0], str) and it[0].startswith("W"))]

# Audit 1: 当日早班≠晚班
for it in all_items:
    date, wd, early, late, backup, off_p, is_we = it
    if early == late:
        errors.append(f"❌ [{date}] 早班=晚班={early}")

# Audit 2: 连续两天早/晚班不重复（含8/31→9/1跨月边界）
for i in range(1, len(all_items)):
    prev = all_items[i-1]
    curr = all_items[i]
    if prev[2] == curr[2]:  # same early person consecutive days
        errors.append(f"❌ [{prev[0]}→{curr[0]}] 连续早班={prev[2]}")
    if prev[3] == curr[3]:  # same late person consecutive days
        errors.append(f"❌ [{prev[0]}→{curr[0]}] 连续晚班={prev[3]}")

# Audit 3: 休假顺序（五=备班人，六=周五早班人，日=周五晚班人）
import re
week_groups = {}
current_wlabel = None
for it in sched:
    if isinstance(it[0], str) and it[0].startswith("W"):
        current_wlabel = it[1]
        week_groups[current_wlabel] = []
    else:
        week_groups[current_wlabel].append(it)

for wdesc, items in week_groups.items():
    fri_item = next((it for it in items if it[1] == "五"), None)
    sat_item = next((it for it in items if it[1] == "六"), None)
    sun_item = next((it for it in items if it[1] == "日"), None)

    m = re.search(r'备班人员[：:]\s*([^（(]+)', wdesc)
    backup_person = m.group(1).strip() if m else "?"

    if fri_item and fri_item[5] != backup_person:
        errors.append(f"❌ [{fri_item[0]}(五)] 休假={fri_item[5]}，应休={backup_person}（备班人）")
    if fri_item and sat_item and sat_item[5] != fri_item[2]:
        errors.append(f"❌ [{sat_item[0]}(六)] 休假={sat_item[5]}，应休={fri_item[2]}（周五早班人）")
    if fri_item and sun_item and sun_item[5] != fri_item[3]:
        errors.append(f"❌ [{sun_item[0]}(日)] 休假={sun_item[5]}，应休={fri_item[3]}（周五晚班人）")

# Audit 4: 备班轮换顺序检查 (二院派遣 → 陈东升 → 孙闪闪)
rotation = ["二院派遣", "陈东升", "孙闪闪"]
backups_in_order = []
for wdesc, items in week_groups.items():
    m = re.search(r'备班人员[：:]\s*([^（(]+)', wdesc)
    if m:
        backups_in_order.append(m.group(1).strip())
for i in range(1, len(backups_in_order)):
    prev_idx = rotation.index(backups_in_order[i-1])
    expected_next = rotation[(prev_idx + 1) % 3]
    if backups_in_order[i] != expected_next:
        errors.append(f"❌ 备班轮换错误：{backups_in_order[i-1]} → {backups_in_order[i]}，应为 {expected_next}")

if errors:
    print(f"\n🔴 发现 {len(errors)} 个问题：")
    for e in errors:
        print(f"  {e}")
else:
    print("✅ 审计全部通过！")
    print("  1. 当日早班≠晚班 ✓")
    print("  2. 连续两天早/晚班不重复（含8/31→9/1跨月边界）✓")
    print("  3. 休假顺序正确（五=备班人，六=周五早班人，日=周五晚班人）✓")
    print(f"  4. 备班轮换顺序正确：{' → '.join(backups_in_order)} ✓")
